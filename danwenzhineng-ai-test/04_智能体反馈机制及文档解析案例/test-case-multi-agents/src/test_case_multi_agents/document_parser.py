"""Lightweight document parsing adapters for the test-case workflow."""

from __future__ import annotations

import csv
import io
import os
import uuid
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any


DEFAULT_MAX_UPLOAD_MB = 20
DEFAULT_MAX_TABLE_ROWS = 200
TEXT_EXTENSIONS = {".md", ".markdown", ".txt"}
CSV_EXTENSIONS = {".csv"}
XLSX_EXTENSIONS = {".xlsx"}
MARKER_EXTENSIONS = {".pdf", ".docx"}


class DocumentParseError(RuntimeError):
    """Raised when an uploaded document cannot be parsed."""


@dataclass(frozen=True)
class ParsedDocument:
    """Normalized parsed document content for downstream agents."""

    document_id: str
    filename: str
    mime_type: str
    parser: str
    content_markdown: str
    metadata: dict[str, Any]
    warnings: list[str]
    parse_quality_score: float | None = None

    def to_dict(self) -> dict[str, Any]:
        """Return a JSON-serializable representation."""
        return asdict(self)


def max_upload_bytes() -> int:
    """Return configured upload limit in bytes."""
    raw_value = os.getenv("DOCUMENT_MAX_UPLOAD_MB", str(DEFAULT_MAX_UPLOAD_MB))
    try:
        max_mb = int(raw_value)
    except ValueError:
        max_mb = DEFAULT_MAX_UPLOAD_MB
    return max_mb * 1024 * 1024


def ensure_allowed_size(content: bytes) -> None:
    """Reject oversized uploads before parsing."""
    limit = max_upload_bytes()
    if len(content) > limit:
        limit_mb = limit // (1024 * 1024)
        raise DocumentParseError(f"Uploaded file is too large. Maximum allowed size is {limit_mb}MB.")


def parse_uploaded_document(filename: str, mime_type: str, content: bytes) -> ParsedDocument:
    """Parse uploaded file bytes into Markdown content."""
    ensure_allowed_size(content)
    suffix = Path(filename).suffix.lower()

    if suffix in TEXT_EXTENSIONS:
        return parse_text_document(filename=filename, mime_type=mime_type, content=content)
    if suffix in CSV_EXTENSIONS:
        return parse_csv_document(filename=filename, mime_type=mime_type, content=content)
    if suffix in XLSX_EXTENSIONS:
        return parse_xlsx_document(filename=filename, mime_type=mime_type, content=content)
    if suffix in MARKER_EXTENSIONS:
        return MarkerParserAdapter().parse(filename=filename, mime_type=mime_type, content=content)

    supported = sorted(TEXT_EXTENSIONS | CSV_EXTENSIONS | XLSX_EXTENSIONS | MARKER_EXTENSIONS)
    raise DocumentParseError(f"Unsupported file type '{suffix or 'unknown'}'. Supported types: {', '.join(supported)}.")


def create_plain_text_document(content: str, filename: str = "pasted-requirement.md") -> ParsedDocument:
    """Wrap manually pasted text as a parsed document for backward compatibility."""
    text = content.strip()
    if len(text) < 10:
        raise DocumentParseError("Requirement content is too short.")
    return ParsedDocument(
        document_id=str(uuid.uuid4()),
        filename=filename,
        mime_type="text/markdown",
        parser="plain_text",
        content_markdown=text,
        metadata={"source": "manual_text", "character_count": len(text)},
        warnings=[],
        parse_quality_score=None,
    )


def parse_text_document(filename: str, mime_type: str, content: bytes) -> ParsedDocument:
    """Parse a UTF-8 text or Markdown document."""
    try:
        text = content.decode("utf-8").strip()
    except UnicodeDecodeError as exc:
        raise DocumentParseError("Text files must be UTF-8 encoded.") from exc

    if not text:
        raise DocumentParseError("Uploaded text document is empty.")

    return ParsedDocument(
        document_id=str(uuid.uuid4()),
        filename=filename,
        mime_type=mime_type or "text/plain",
        parser="local_text",
        content_markdown=text,
        metadata={"character_count": len(text), "byte_count": len(content)},
        warnings=[],
    )


def parse_csv_document(filename: str, mime_type: str, content: bytes) -> ParsedDocument:
    """Parse a CSV document and render it as a Markdown table."""
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentParseError("CSV files must be UTF-8 encoded.") from exc

    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        raise DocumentParseError("Uploaded CSV document is empty.")

    warnings: list[str] = []
    rendered_rows = rows[:DEFAULT_MAX_TABLE_ROWS]
    if len(rows) > DEFAULT_MAX_TABLE_ROWS:
        warnings.append(f"CSV contains {len(rows)} rows; only the first {DEFAULT_MAX_TABLE_ROWS} rows were parsed.")

    table = rows_to_markdown_table(rendered_rows)
    return ParsedDocument(
        document_id=str(uuid.uuid4()),
        filename=filename,
        mime_type=mime_type or "text/csv",
        parser="local_csv",
        content_markdown=f"# {filename}\n\n{table}",
        metadata={"row_count": len(rows), "column_count": max(len(row) for row in rows), "byte_count": len(content)},
        warnings=warnings,
    )


def parse_xlsx_document(filename: str, mime_type: str, content: bytes) -> ParsedDocument:
    """Parse an XLSX workbook and render each sheet as a Markdown table."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise DocumentParseError("XLSX parsing requires the optional dependency 'openpyxl'.") from exc

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    warnings: list[str] = []
    sections: list[str] = [f"# {filename}"]
    sheet_metadata: list[dict[str, Any]] = []

    for worksheet in workbook.worksheets:
        rows = [
            [cell_to_text(value) for value in row]
            for row in worksheet.iter_rows(values_only=True)
            if any(value is not None and str(value).strip() for value in row)
        ]
        if not rows:
            sections.append(f"\n## Sheet: {worksheet.title}\n\n_No data found._")
            sheet_metadata.append({"name": worksheet.title, "row_count": 0, "column_count": 0})
            continue

        rendered_rows = rows[:DEFAULT_MAX_TABLE_ROWS]
        if len(rows) > DEFAULT_MAX_TABLE_ROWS:
            warnings.append(
                f"Sheet '{worksheet.title}' contains {len(rows)} rows; only the first "
                f"{DEFAULT_MAX_TABLE_ROWS} rows were parsed."
            )

        sections.append(f"\n## Sheet: {worksheet.title}\n\n{rows_to_markdown_table(rendered_rows)}")
        sheet_metadata.append(
            {"name": worksheet.title, "row_count": len(rows), "column_count": max(len(row) for row in rows)}
        )

    return ParsedDocument(
        document_id=str(uuid.uuid4()),
        filename=filename,
        mime_type=mime_type or "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        parser="local_xlsx",
        content_markdown="\n".join(sections).strip(),
        metadata={"sheet_count": len(workbook.worksheets), "sheets": sheet_metadata, "byte_count": len(content)},
        warnings=warnings,
    )


def rows_to_markdown_table(rows: list[list[str]]) -> str:
    """Render rectangular-ish rows as a Markdown table."""
    column_count = max(len(row) for row in rows)
    normalized = [row + [""] * (column_count - len(row)) for row in rows]
    header = [value or f"Column {index + 1}" for index, value in enumerate(normalized[0])]
    body = normalized[1:]

    lines = [
        "| " + " | ".join(escape_markdown_cell(value) for value in header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    for row in body:
        lines.append("| " + " | ".join(escape_markdown_cell(value) for value in row) + " |")
    return "\n".join(lines)


def cell_to_text(value: object) -> str:
    """Convert spreadsheet cell values into compact text."""
    if value is None:
        return ""
    return str(value).strip()


def escape_markdown_cell(value: str) -> str:
    """Escape Markdown table delimiters and normalize whitespace."""
    return " ".join(value.replace("|", "\\|").split())


class MarkerParserAdapter:
    """Reserved adapter for Marker/Datalab-based complex document parsing."""

    def parse(self, filename: str, mime_type: str, content: bytes) -> ParsedDocument:
        """Return a clear setup error until a concrete Marker backend is configured."""
        _ = content
        parser_mode = os.getenv("DOCUMENT_PARSER_MODE", "local_first")
        datalab_key = os.getenv("DATALAB_API_KEY")
        try:
            import marker  # type: ignore  # noqa: F401

            marker_installed = True
        except ImportError:
            marker_installed = False

        if not marker_installed and not datalab_key:
            raise DocumentParseError(
                "PDF/DOCX parsing is reserved for Marker or Datalab, but neither local marker nor "
                "DATALAB_API_KEY is configured. Use md/txt/csv/xlsx for local parsing, or enable a "
                "Marker/Datalab adapter."
            )

        raise DocumentParseError(
            "PDF/DOCX Marker parsing adapter is configured as a placeholder in this version. "
            f"Detected parser mode: {parser_mode}. Wire local marker or Datalab Convert API in "
            "MarkerParserAdapter.parse before enabling complex document parsing."
        )
