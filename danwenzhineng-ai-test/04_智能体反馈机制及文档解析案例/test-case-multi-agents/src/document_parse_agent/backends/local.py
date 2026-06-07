"""Lightweight local parsers for text, CSV, and XLSX documents."""

from __future__ import annotations

import csv
import io
import uuid
from pathlib import Path
from typing import Any

from document_parse_agent.errors import DocumentParseError
from document_parse_agent.models import ParseOptions, ParseResult, ParseWarning, ParserBackend


TEXT_EXTENSIONS = {".md", ".markdown", ".txt"}
CSV_EXTENSIONS = {".csv"}
XLSX_EXTENSIONS = {".xlsx"}
LOCAL_EXTENSIONS = TEXT_EXTENSIONS | CSV_EXTENSIONS | XLSX_EXTENSIONS


def supports_local_file(filename: str) -> bool:
    """Return whether the local backend supports a filename."""
    return Path(filename).suffix.lower() in LOCAL_EXTENSIONS


def parse_text(content: str, filename: str, options: ParseOptions | None = None) -> ParseResult:
    """Wrap manually supplied text as a parse result."""
    _ = options
    text = content.strip()
    if len(text) < 10:
        raise DocumentParseError("Requirement content is too short.")

    return ParseResult(
        document_id=str(uuid.uuid4()),
        filename=filename,
        mime_type="text/markdown",
        backend=ParserBackend.PLAIN_TEXT.value,
        output_format="markdown",
        content_markdown=text,
        content_raw=text,
        metadata={"source": "manual_text", "character_count": len(text)},
    )


def parse_text_bytes(filename: str, mime_type: str, content: bytes, options: ParseOptions) -> ParseResult:
    """Parse a UTF-8 text or Markdown document."""
    try:
        text = content.decode("utf-8").strip()
    except UnicodeDecodeError as exc:
        raise DocumentParseError("Text files must be UTF-8 encoded.") from exc

    if not text:
        raise DocumentParseError("Uploaded text document is empty.")

    return ParseResult(
        document_id=str(uuid.uuid4()),
        filename=filename,
        mime_type=mime_type or "text/plain",
        backend=ParserBackend.LOCAL_TEXT.value,
        output_format=options.output_format,
        content_markdown=text,
        content_raw=text,
        metadata={"character_count": len(text), "byte_count": len(content)},
    )


def parse_csv_bytes(filename: str, mime_type: str, content: bytes, options: ParseOptions) -> ParseResult:
    """Parse a CSV document and render it as a Markdown table."""
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentParseError("CSV files must be UTF-8 encoded.") from exc

    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        raise DocumentParseError("Uploaded CSV document is empty.")

    warnings: list[ParseWarning] = []
    rendered_rows = rows[: options.max_table_rows]
    if len(rows) > options.max_table_rows:
        warnings.append(
            ParseWarning(
                code="row_limit",
                message=f"CSV contains {len(rows)} rows; only the first {options.max_table_rows} rows were parsed.",
                details={"row_count": len(rows), "rendered_rows": options.max_table_rows},
            )
        )

    table = rows_to_markdown_table(rendered_rows)
    return ParseResult(
        document_id=str(uuid.uuid4()),
        filename=filename,
        mime_type=mime_type or "text/csv",
        backend=ParserBackend.LOCAL_CSV.value,
        output_format=options.output_format,
        content_markdown=f"# {filename}\n\n{table}",
        content_raw={"rows": rows},
        metadata={"row_count": len(rows), "column_count": max(len(row) for row in rows), "byte_count": len(content)},
        warnings=warnings,
    )


def parse_xlsx_bytes(filename: str, mime_type: str, content: bytes, options: ParseOptions) -> ParseResult:
    """Parse an XLSX workbook and render each sheet as a Markdown table."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise DocumentParseError("XLSX parsing requires the optional dependency 'openpyxl'.") from exc

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    warnings: list[ParseWarning] = []
    sections: list[str] = [f"# {filename}"]
    sheet_metadata: list[dict[str, Any]] = []

    for worksheet in workbook.worksheets:
        rows = [
            [cell_to_text(value) for value in row]
            for row in worksheet.iter_rows(values_only=True)
            if any(value is not None and str(value).strip() for value in row)
        ]
        if not rows:
            sections.append(f"\n## Sheet: {worksheet.title}\n\n_No data found._")
            sheet_metadata.append({"name": worksheet.title, "row_count": 0, "column_count": 0})
            continue

        rendered_rows = rows[: options.max_table_rows]
        if len(rows) > options.max_table_rows:
            warnings.append(
                ParseWarning(
                    code="row_limit",
                    message=(
                        f"Sheet '{worksheet.title}' contains {len(rows)} rows; only the first "
                        f"{options.max_table_rows} rows were parsed."
                    ),
                    details={"sheet": worksheet.title, "row_count": len(rows), "rendered_rows": options.max_table_rows},
                )
            )

        sections.append(f"\n## Sheet: {worksheet.title}\n\n{rows_to_markdown_table(rendered_rows)}")
        sheet_metadata.append(
            {"name": worksheet.title, "row_count": len(rows), "column_count": max(len(row) for row in rows)}
        )

    return ParseResult(
        document_id=str(uuid.uuid4()),
        filename=filename,
        mime_type=mime_type or "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        backend=ParserBackend.LOCAL_XLSX.value,
        output_format=options.output_format,
        content_markdown="\n".join(sections).strip(),
        content_raw=None,
        metadata={"sheet_count": len(workbook.worksheets), "sheets": sheet_metadata, "byte_count": len(content)},
        warnings=warnings,
    )


def parse_local_bytes(filename: str, mime_type: str, content: bytes, options: ParseOptions) -> ParseResult:
    """Parse local-supported file bytes into a parse result."""
    suffix = Path(filename).suffix.lower()
    if suffix in TEXT_EXTENSIONS:
        return parse_text_bytes(filename=filename, mime_type=mime_type, content=content, options=options)
    if suffix in CSV_EXTENSIONS:
        return parse_csv_bytes(filename=filename, mime_type=mime_type, content=content, options=options)
    if suffix in XLSX_EXTENSIONS:
        return parse_xlsx_bytes(filename=filename, mime_type=mime_type, content=content, options=options)
    raise DocumentParseError(f"Local parser does not support file type '{suffix or 'unknown'}'.")


def rows_to_markdown_table(rows: list[list[str]]) -> str:
    """Render rectangular-ish rows as a Markdown table."""
    column_count = max(len(row) for row in rows)
    normalized = [row + [""] * (column_count - len(row)) for row in rows]
    header = [value or f"Column {index + 1}" for index, value in enumerate(normalized[0])]
    body = normalized[1:]

    lines = [
        "| " + " | ".join(escape_markdown_cell(value) for value in header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    for row in body:
        lines.append("| " + " | ".join(escape_markdown_cell(value) for value in row) + " |")
    return "\n".join(lines)


def cell_to_text(value: object) -> str:
    """Convert spreadsheet cell values into compact text."""
    if value is None:
        return ""
    return str(value).strip()


def escape_markdown_cell(value: str) -> str:
    """Escape Markdown table delimiters and normalize whitespace."""
    return " ".join(value.replace("|", "\\|").split())
